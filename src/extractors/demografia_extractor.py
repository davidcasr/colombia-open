"""Extractor de proyecciones de población municipal, fuente DANE.

Fuente primaria: descarga directa (no API) del archivo oficial de
Proyecciones de Población y Estudios Demográficos (PPED), hoja
"PobMunicipalxÁrea" — población municipal por área geográfica
(cabecera / centros poblados y rural disperso / total), años 2018-2042.
Confirmado en vivo el 2026-10-18: 84225 filas, 1123 municipios, 33
departamentos, 25 años, 3 categorías de área geográfica.

https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/Municipal/PPED-AreaMun-2018-2042_VP.xlsx

Nota de licencia: este archivo se descarga de dane.gov.co directamente (NO
de datos.gov.co/Socrata, que es donde `divipola` sí declara CC BY-SA 4.0
explícito). Los términos de uso del DANE autorizan "uso, aprovechamiento,
transformación y análisis" con atribución obligatoria, pero no mencionan
redistribución ni ninguna licencia Creative Commons explícita — ver
`docs/licenses/data_licenses.md`. Se publica como `public-api-review-terms`
(dato estatal público, Ley 1712/2014 de apertura por defecto), no como
`open-attribution`.

Nota de alcance: 1 de los 1123 códigos de municipio de esta fuente
(`94663`, Mapiripana (ANM)) no existe en el `municipios` vigente de
`divipola` — el municipio de Barrancominas (94343) absorbió esa área no
municipalizada desde el 1 de diciembre de 2019 (ver nota del propio
archivo DANE), pero esta serie histórica conserva sus filas 2018-2042 bajo
el código antiguo. Es una discrepancia documentada por el DANE entre dos
productos suyos con distinto ciclo de actualización, no un error de
extracción — el validador la reporta como warning, no como error (ver
`src/validation/_demografia.py`).

Se invoca siempre como módulo, nunca como script suelto:
`python -m src.extractors.demografia_extractor` desde la raíz del repo (así
`src` se resuelve como paquete). Ver `.claude/rules/extractors.md` §6.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import polars as pl
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from requests import RequestException

from src.extractors.base import (
    RAW_DIR,
    ROOT_DIR,
    STAGING_DIR,
    BaseExtractor,
    ExtractionResult,
    ReusePolicy,
    compute_sha256,
    ensure_staging_directories,
    write_raw_snapshot_atomic,
    write_staging_csv_atomic,
)
from src.extractors.http_utils import get_bytes

DATASET_NAME = "demografia"
SOURCE_URL = (
    "https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/"
    "Municipal/PPED-AreaMun-2018-2042_VP.xlsx"
)
SOURCE_NAME = "DANE — Proyecciones de población municipal por área geográfica 2018-2042"
SHEET_NAME = "PobMunicipalxÁrea"

# La fuente conocida tiene 84225 filas (2026-10-18). Un umbral generoso
# detecta truncamiento severo (ej. el archivo cambió de formato o la hoja
# llegó vacía) sin reventar por variaciones normales año a año.
EXPECTED_MIN_ROWS = 70_000

REUSE_POLICY = ReusePolicy(
    status="public-api-review-terms",
    license="Sin licencia abierta explícita — términos de uso del DANE",
    license_url=(
        "https://www.dane.gov.co/index.php/servicios-al-ciudadano/tramites/"
        "transparencia-y-acceso-a-la-informacion-publica/terminos-y-condiciones"
    ),
    attribution_required=True,
    redistribution_ok=True,
    summary=(
        "Dato estatal público descargado directamente de dane.gov.co. Los "
        "términos de uso del DANE autorizan uso, transformación y análisis "
        "con atribución obligatoria ('Fuente: Departamento Administrativo "
        "Nacional de Estadística: www.dane.gov.co'), pero no declaran "
        "Creative Commons ni mencionan redistribución explícitamente. Se "
        "publica bajo la presunción de apertura de la Ley 1712/2014 de "
        "Transparencia (información pública estadística, sin restricción "
        "de reuso declarada). Ver docs/licenses/data_licenses.md."
    ),
)

# Qué cae a fallback: problemas de la FUENTE (red caída, HTTP 4xx/5xx tras
# agotar reintentos, el archivo descargado no es un .xlsx válido —
# BadZipFile/InvalidFileException—, o la hoja/columnas esperadas no están
# donde se esperaba — KeyError/ValueError/TypeError). Un bug de
# programación real (AttributeError, NameError, etc.) NO está acá a
# propósito: debe reventar fuerte en vez de disfrazarse de "fuente caída".
# Ver `.claude/rules/extractors.md` §1.
EXPECTED_FETCH_FAILURES: tuple[type[Exception], ...] = (
    RequestException,
    BadZipFile,
    InvalidFileException,
    KeyError,
    ValueError,
    TypeError,
)

# Fallback mínimo para desarrollo offline / CI sin red. NUNCA se publica:
# la CI de publicación rechaza cualquier dataset con source_mode="fallback".
# Filas reales observadas en la fuente (2026-10-18), recortadas a 3
# municipios x 3 áreas geográficas para cubrir un año temprano y uno tardío.
DEMOGRAFIA_FALLBACK_DATA: list[tuple[Any, ...]] = [
    ("05", "Antioquia", "05001", "Medellín", 2018, "Cabecera Municipal", 2382405),
    ("05", "Antioquia", "05001", "Medellín", 2018, "Centros Poblados y Rural Disperso", 44728),
    ("05", "Antioquia", "05001", "Medellín", 2018, "Total", 2427133),
    ("05", "Antioquia", "05002", "Abejorral", 2018, "Cabecera Municipal", 7669),
    ("05", "Antioquia", "05002", "Abejorral", 2018, "Centros Poblados y Rural Disperso", 12699),
    ("05", "Antioquia", "05002", "Abejorral", 2018, "Total", 20368),
    ("99", "Vichada", "99773", "Cumaribo", 2042, "Cabecera Municipal", 4461),
    ("99", "Vichada", "99773", "Cumaribo", 2042, "Centros Poblados y Rural Disperso", 176148),
    ("99", "Vichada", "99773", "Cumaribo", 2042, "Total", 180609),
]

_SCHEMA = {
    "cod_dpto": pl.Utf8,
    "nombre_dpto": pl.Utf8,
    "cod_mpio": pl.Utf8,
    "nombre_mpio": pl.Utf8,
    "anio": pl.Int64,
    "area_geografica": pl.Utf8,
    "poblacion": pl.Int64,
}


def _is_data_row(row: tuple[Any, ...]) -> bool:
    """Distingue una fila de dato real de encabezados/separadores/notas al
    pie: una fila de dato siempre tiene `cod_dpto` como texto de 2 dígitos
    numéricos y `anio` como entero. No depende de números de fila fijos —
    el archivo DANE intercala filas vacías y notas sin aviso."""
    cod_dpto, anio = row[0], row[4]
    return (
        isinstance(cod_dpto, str)
        and cod_dpto.isdigit()
        and len(cod_dpto) == 2
        and isinstance(anio, int)
    )


def _normalize(rows: list[tuple[Any, ...]]) -> pl.DataFrame:
    records = [
        {
            "cod_dpto": str(row[0]).zfill(2),
            "nombre_dpto": row[1],
            "cod_mpio": str(row[2]).zfill(5),
            "nombre_mpio": row[3],
            "anio": int(row[4]),
            "area_geografica": row[5],
            "poblacion": int(row[6]),
        }
        for row in rows
    ]
    df = pl.DataFrame(records, schema=_SCHEMA)
    return df.unique(subset=["cod_mpio", "anio", "area_geografica"], keep="first").sort(
        ["cod_mpio", "anio", "area_geografica"]
    )


def _parse_workbook_rows(xlsx_bytes: bytes) -> list[tuple[Any, ...]]:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    worksheet = workbook[SHEET_NAME]
    return [row for row in worksheet.iter_rows(values_only=True) if _is_data_row(row)]


class DemografiaExtractor(BaseExtractor):
    @property
    def dataset_name(self) -> str:
        return DATASET_NAME

    def fetch(self, **kwargs: Any) -> ExtractionResult:
        notes: list[str] = []
        try:
            xlsx_bytes = get_bytes(SOURCE_URL)
            rows = _parse_workbook_rows(xlsx_bytes)
            if len(rows) < EXPECTED_MIN_ROWS:
                notes.append(
                    f"La fuente devolvió solo {len(rows)} filas, menos de las "
                    f"{EXPECTED_MIN_ROWS} esperadas como mínimo."
                )
            df = _normalize(rows)

            ensure_staging_directories()
            timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
            raw_path = RAW_DIR / f"demografia_{timestamp}.xlsx"
            write_raw_snapshot_atomic(raw_path, xlsx_bytes)

            return ExtractionResult(
                dataset=DATASET_NAME,
                dataframe=df,
                source_name=SOURCE_NAME,
                source_url=SOURCE_URL,
                source_mode="live",
                source_detail=f"Excel directo, hoja '{SHEET_NAME}'",
                reuse_policy=REUSE_POLICY,
                raw_snapshot_path=str(raw_path.relative_to(ROOT_DIR)),
                snapshot_hash=compute_sha256(xlsx_bytes),
                notes=notes,
            )
        except EXPECTED_FETCH_FAILURES as exc:
            df = _normalize(DEMOGRAFIA_FALLBACK_DATA)
            return ExtractionResult(
                dataset=DATASET_NAME,
                dataframe=df,
                source_name=SOURCE_NAME,
                source_url=SOURCE_URL,
                source_mode="fallback",
                source_detail=f"fallback embebido tras fallo de fetch en vivo: {exc!r}",
                reuse_policy=REUSE_POLICY,
                notes=[
                    "source_mode=fallback: dataset de muestra para desarrollo offline, "
                    "NO se publica."
                ],
            )

    def write_staging(self, result: ExtractionResult) -> Path:
        ensure_staging_directories()
        csv_path = STAGING_DIR / "demografia.csv"
        metadata_path = STAGING_DIR / "demografia.metadata.json"
        return write_staging_csv_atomic(
            result.dataframe, csv_path, metadata_path, result.to_staging_metadata()
        )


if __name__ == "__main__":
    extractor = DemografiaExtractor()
    outcome = extractor.run()
    print(f"[{DATASET_NAME}] source_mode={outcome.source_mode} rows={outcome.dataframe.height}")
